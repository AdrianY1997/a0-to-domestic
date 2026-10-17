#!/usr/bin/env python3
"""Deterministic fixture generator (SPEC 15.1, 15.6).

Usage: python make_fixtures.py [--out fixtures] [--large]
No timestamps / random values: outputs are byte-identical between runs.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import struct
import zipfile
import zlib
from pathlib import Path

import numpy as np
import pymupdf
from PIL import Image

MM = 72.0 / 25.4
FIXED_DATE = "D:20200101000000Z"
SIZES = {"A0": (841, 1189), "A1": (594, 841), "A3": (297, 420)}
MANIFEST: dict[str, dict] = {}
FIXED_ID = "a0a0a0a0a0a0a0a0a0a0a0a0a0a0a0a0"


def pt(mm: float) -> float:
    return mm * MM


def note(rel, fmt, pages, size_mm, notes, expect="accept"):
    MANIFEST[rel] = {"format": fmt, "pages": pages, "size_mm": size_mm,
                     "expect": expect, "notes": notes}


# ---------------------------------------------------------------- PDF helpers
def new_doc() -> pymupdf.Document:
    doc = pymupdf.open()
    doc.set_metadata({"producer": "a0td-fixtures", "creator": "a0td-fixtures",
                      "title": "fixture", "author": "", "subject": "", "keywords": "",
                      "creationDate": FIXED_DATE, "modDate": FIXED_DATE})
    return doc


def finalize(data: bytes) -> bytes:
    """Replace trailer /ID with a fixed value (unencrypted files only)."""
    rep = ("/ID[<%s><%s>]" % (FIXED_ID, FIXED_ID)).encode()
    return re.sub(rb"/ID\s*\[\s*<[0-9a-fA-F]+>\s*<[0-9a-fA-F]+>\s*\]", lambda m: rep, data)


def save(doc, path: Path, **kw):
    kw.setdefault("garbage", 0)
    kw.setdefault("deflate", True)
    data = doc.tobytes(**kw)
    path.write_bytes(finalize(data))


def draw_grid(page, w_mm, h_mm, y1_mm=None):
    """10 mm grid, labeled 100 mm squares, coordinate text every 100 mm, hairlines."""
    y1_mm = h_mm if y1_mm is None else y1_mm
    sh = page.new_shape()
    for x in range(0, int(w_mm) + 1, 10):
        major = x % 100 == 0
        sh.draw_line((pt(x), 0), (pt(x), pt(y1_mm)))
        sh.finish(color=(0, 0, 0) if major else (0.6, 0.6, 0.9),
                  width=1.0 if major else 0.25, closePath=False)
    for y in range(0, int(y1_mm) + 1, 10):
        major = y % 100 == 0
        sh.draw_line((0, pt(y)), (pt(w_mm), pt(y)))
        sh.finish(color=(0, 0, 0) if major else (0.6, 0.6, 0.9),
                  width=1.0 if major else 0.25, closePath=False)
    sh.draw_line((0, 0), (pt(w_mm), pt(y1_mm)))
    sh.finish(color=(1, 0, 0), width=0.01)
    sh.draw_line((pt(w_mm), 0), (0, pt(y1_mm)))
    sh.finish(color=(0, 0.5, 0), width=0)
    sh.commit()
    for xi in range(0, int(w_mm), 100):
        for yi in range(0, int(y1_mm), 100):
            page.insert_text((pt(xi + 3), pt(yi + 10)), "x=%d y=%d mm" % (xi, yi),
                             fontsize=14, fontname="helv")
            page.insert_text((pt(xi + 3), pt(yi + 50)), "[%d,%d]" % (xi // 100, yi // 100),
                             fontsize=48, fontname="helv", color=(0.7, 0, 0))


def grid_pdf(path, w_mm, h_mm, y1_mm=None, rotate=0):
    doc = new_doc()
    page = doc.new_page(width=pt(w_mm), height=pt(h_mm))
    draw_grid(page, w_mm, h_mm, y1_mm)
    if rotate:
        page.set_rotation(rotate)
    save(doc, path)


def raw_pdf(objs: list[bytes], root: int = 1) -> bytes:
    out = bytearray(b"%PDF-1.7\n")
    offs = []
    for i, body in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for o in offs:
        out += b"%010d 00000 n \n" % o
    out += b"trailer\n<</Size %d/Root %d 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, root, xref)
    return bytes(out)



_PAD = bytes.fromhex("28BF4E5E4E758A4164004E56FFFA01082E2E00B6D0683E802F0CA9FE6453697A")


def _rc4(key: bytes, data: bytes) -> bytes:
    S = list(range(256))
    j = 0
    for i in range(256):
        j = (j + S[i] + key[i % len(key)]) & 255
        S[i], S[j] = S[j], S[i]
    i = j = 0
    out = bytearray()
    for c in data:
        i = (i + 1) & 255
        j = (j + S[i]) & 255
        S[i], S[j] = S[j], S[i]
        out.append(c ^ S[(S[i] + S[j]) & 255])
    return bytes(out)


def encrypted_pdf(user_pw: str, owner_pw: str, perms: int, title: str) -> bytes:
    """Standard security handler R3 (RC4-128) written by hand with a fixed /ID (deterministic)."""
    md5 = lambda b: hashlib.md5(b).digest()
    pad = lambda pw: (pw.encode("latin-1") + _PAD)[:32]
    ident = bytes.fromhex(FIXED_ID)
    h = md5(pad(owner_pw))
    for _ in range(50):
        h = md5(h)
    okey = h[:16]
    o = _rc4(okey, pad(user_pw))
    for i in range(1, 20):
        o = _rc4(bytes(b ^ i for b in okey), o)
    h = md5(pad(user_pw) + o + struct.pack("<i", perms) + ident)
    for _ in range(50):
        h = md5(h[:16])
    key = h[:16]
    u = _rc4(key, md5(_PAD + ident))
    for i in range(1, 20):
        u = _rc4(bytes(b ^ i for b in key), u)
    u += b"\x00" * 16
    lines = ["0.5 w 0 0 0 RG"]
    for x in range(0, 840, 100):
        lines.append("%d 0 m %d 1190 l S" % (x, x))
        lines.append("0 %d m 840 %d l S" % (x, x))
    lines.append("BT /F1 40 Tf 40 600 Td (%s) Tj ET" % title)
    content = ("\n".join(lines) + "\n").encode()

    def okey_for(n):
        return md5(key + struct.pack("<I", n)[:3] + b"\x00\x00")[:16]
    hexs = lambda b: b"<" + b.hex().encode() + b">"
    bodies = {
        1: b"<</Type/Catalog/Pages 2 0 R>>",
        2: b"<</Type/Pages/Count 1/Kids[3 0 R]>>",
        3: b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 842 1190]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        5: b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
        6: b"<</Filter/Standard/V 2/R 3/Length 128/P %d/O %s/U %s>>" % (perms, hexs(o), hexs(u)),
    }
    enc = _rc4(okey_for(4), content)
    bodies[4] = b"<</Length %d>>\nstream\n" % len(enc) + enc + b"\nendstream"
    out = bytearray(b"%PDF-1.6\n")
    offs = []
    for n in range(1, 7):
        offs.append(len(out))
        out += b"%d 0 obj\n" % n + bodies[n] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 7\n0000000000 65535 f \n"
    for off in offs:
        out += b"%010d 00000 n \n" % off
    out += (b"trailer\n<</Size 7/Root 1 0 R/Encrypt 6 0 R/ID[%s%s]>>\nstartxref\n%d\n%%%%EOF\n"
            % (hexs(ident), hexs(ident), xref))
    return bytes(out)


# ---------------------------------------------------------------- F1 PDFs
def make_pdfs(out: Path):
    grid_pdf(out / "a0_vector_grid.pdf", 841, 1189)
    note("a0_vector_grid.pdf", "pdf", 1, [841, 1189], "10mm grid, labeled 100mm squares, hairlines")
    grid_pdf(out / "a0_top_half_only.pdf", 841, 1189, y1_mm=594.5)
    note("a0_top_half_only.pdf", "pdf", 1, [841, 1189], "content only in top half (blank bottom)")
    grid_pdf(out / "a1_landscape.pdf", 841, 594)
    note("a1_landscape.pdf", "pdf", 1, [841, 594], "A1 landscape")
    grid_pdf(out / "rotated_90.pdf", 297, 420, rotate=90)
    note("rotated_90.pdf", "pdf", 1, [420, 297], "A3 portrait mediabox with /Rotate 90 (effective landscape)")

    doc = new_doc()
    page = doc.new_page(width=pt(420), height=pt(594))
    draw_grid(page, 420, 594)
    mb = page.mediabox
    page.set_cropbox(pymupdf.Rect(pt(20), pt(20), mb.x1 - pt(20), mb.y1 - pt(20)))
    page.set_trimbox(pymupdf.Rect(pt(30), pt(30), mb.x1 - pt(30), mb.y1 - pt(30)))
    page.set_bleedbox(pymupdf.Rect(pt(25), pt(25), mb.x1 - pt(25), mb.y1 - pt(25)))
    save(doc, out / "boxes.pdf")
    note("boxes.pdf", "pdf", 1, [380, 554],
         "MediaBox 420x594mm; CropBox inset 20mm (effective size); TrimBox inset 30mm; BleedBox inset 25mm")

    (out / "encrypted.pdf").write_bytes(
        encrypted_pdf("secret", "owner-secret", -4, "Encrypted A3 (user pw: secret)"))
    note("encrypted.pdf", "pdf", 1, [297, 420], "user password 'secret'", expect="password_required")
    (out / "encrypted_noprint.pdf").write_bytes(
        encrypted_pdf("", "owner-secret", -3904, "No-print A3 (owner pw: owner-secret)"))
    note("encrypted_noprint.pdf", "pdf", 1, [297, 420],
         "empty user pw, owner pw 'owner-secret', printing not permitted", expect="permission_denied")

    doc = new_doc()
    page = doc.new_page(width=pt(297), height=pt(420))
    draw_grid(page, 297, 420)
    a = page.add_stamp_annot(pymupdf.Rect(pt(20), pt(20), pt(80), pt(50)), stamp=0)
    a.set_info(creationDate=FIXED_DATE, modDate=FIXED_DATE)
    a = page.add_freetext_annot(pymupdf.Rect(pt(100), pt(20), pt(200), pt(40)), "Freetext note", fontsize=12)
    a.set_info(creationDate=FIXED_DATE, modDate=FIXED_DATE)
    a = page.add_ink_annot([[(pt(30), pt(100)), (pt(60), pt(130)), (pt(90), pt(100))]])
    a.set_colors(stroke=(1, 0, 0))
    a.set_info(creationDate=FIXED_DATE, modDate=FIXED_DATE)
    w = pymupdf.Widget()
    w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    w.field_name = "name"
    w.field_value = "filled value"
    w.rect = pymupdf.Rect(pt(20), pt(200), pt(120), pt(215))
    w.fill_color = (0.9, 0.9, 0.5)
    page.add_widget(w)
    save(doc, out / "annotations.pdf")
    note("annotations.pdf", "pdf", 1, [297, 420], "stamp, freetext, ink annotations + filled text form field")

    doc = new_doc()
    page = doc.new_page(width=pt(841), height=pt(594))
    for i, name in enumerate(["Talla 38", "Talla 40", "Talla 42"]):
        oc = doc.add_ocg(name, on=True)
        s = 1 + 0.05 * i
        sh = page.new_shape()
        cx, cy = pt(200 + 220 * i), pt(300)
        pts = [(cx - pt(80 * s), cy - pt(150 * s)), (cx + pt(80 * s), cy - pt(150 * s)),
               (cx + pt(100 * s), cy + pt(150 * s)), (cx - pt(100 * s), cy + pt(150 * s))]
        sh.draw_polyline(pts + [pts[0]])
        sh.draw_circle((cx, cy - pt(150 * s)), pt(25 * s))
        sh.finish(color=(0, 0, 0.8), fill=(0.85, 0.9, 1), width=1.5, oc=oc)
        sh.commit()
        page.insert_text((cx - 40, cy), name, fontsize=24, oc=oc)
    save(doc, out / "layers_ocg.pdf")
    note("layers_ocg.pdf", "pdf", 1, [841, 594], "pattern-like shapes in OCG layers 'Talla 38','Talla 40','Talla 42'")

    doc = new_doc()
    for k in ("A0", "A1", "A3"):
        w_, h_ = SIZES[k]
        draw_grid(doc.new_page(width=pt(w_), height=pt(h_)), w_, h_)
    save(doc, out / "multipage_mixed.pdf")
    note("multipage_mixed.pdf", "pdf", 3, [841, 1189], "pages: A0, A1, A3 (all portrait)")

    doc = new_doc()
    for _ in range(3):
        draw_grid(doc.new_page(width=pt(297), height=pt(420)), 297, 420)
    data = finalize(doc.tobytes(garbage=0, deflate=False))
    cut = data[: int(len(data) * 0.6)]
    (out / "corrupt.pdf").write_bytes(cut)
    note("corrupt.pdf", "pdf", None, None, "truncated at 60%, no xref/trailer", expect="reject_or_repair")


# ---------------------------------------------------------------- raster
def make_rasters(out: Path, large: bool):
    W, H = 6622, 9362
    img = np.ones((H, W), dtype=bool)  # True = white
    step = 79  # ~10 mm @200dpi
    img[::step, :] = False
    img[:, ::step] = False
    big = 790
    for k in range(0, H, big):
        img[k:k + 3, :] = False
    for k in range(0, W, big):
        img[:, k:k + 3] = False
    for gy in range(0, H // big + 1):
        for gx in range(0, W // big + 1):
            y, x = gy * big + 12, gx * big + 12
            img[y:y + 40, x:x + 8 * (gx + 1)] = False
            img[y + 50:y + 90, x:x + 8 * (gy + 1)] = False
    Image.fromarray(img).convert("1").save(out / "scan_a0_200dpi_g4.tif", compression="group4", dpi=(200, 200))
    note("scan_a0_200dpi_g4.tif", "tiff", 1, [840.9, 1188.9], "1-bit CCITT G4, 200 dpi, 6622x9362")

    rng = np.random.default_rng(12345)
    h, w = 5657, 4000
    yy = np.linspace(0, 255, h, dtype=np.float32)[:, None]
    xx = np.linspace(0, 255, w, dtype=np.float32)[None, :]
    noise = rng.integers(0, 24, size=(h // 8 + 1, w // 8 + 1), dtype=np.uint8)
    noise = np.kron(noise, np.ones((8, 8), dtype=np.uint8))[:h, :w].astype(np.float32)
    rgb = np.stack([yy + 0 * xx, xx + 0 * yy, (xx + yy) / 2], axis=-1) * 0.9 + noise[..., None]
    Image.fromarray(np.clip(rgb, 0, 255).astype(np.uint8), "RGB").save(out / "photo_nodpi.jpg", quality=80)
    note("photo_nodpi.jpg", "jpeg", 1, None, "4000x5657 px, no DPI metadata (size must be asked from the user)")

    a = np.zeros((300, 400, 4), dtype=np.uint8)
    a[..., 0] = np.linspace(0, 255, 400, dtype=np.uint8)[None, :]
    a[..., 2] = 200
    a[..., 3] = np.linspace(0, 255, 300, dtype=np.uint8)[:, None]
    Image.fromarray(a, "RGBA").save(out / "transparent.png", dpi=(150, 150))
    note("transparent.png", "png", 1, [67.7, 50.8], "RGBA gradient alpha, 400x300 @150dpi")

    c = np.zeros((300, 400, 4), dtype=np.uint8)
    c[..., 0] = np.linspace(0, 255, 400, dtype=np.uint8)[None, :]
    c[..., 1] = np.linspace(0, 255, 300, dtype=np.uint8)[:, None]
    c[..., 3] = 20
    Image.fromarray(c, "CMYK").save(out / "cmyk.jpg", quality=90, dpi=(150, 150))
    note("cmyk.jpg", "jpeg", 1, [67.7, 50.8], "CMYK JPEG 400x300 @150dpi")

    pages = []
    for i in range(3):
        arr = np.full((400, 300, 3), 255, dtype=np.uint8)
        arr[::20, :] = (0, 0, 0)
        arr[:, ::20] = (0, 0, 0)
        arr[100:200, 50 + 50 * i:150 + 50 * i] = (255 * (i == 0), 255 * (i == 1), 255 * (i == 2))
        pages.append(Image.fromarray(arr, "RGB"))
    pages[0].save(out / "multipage.tif", save_all=True, append_images=pages[1:], dpi=(100, 100))
    note("multipage.tif", "tiff", 3, [76.2, 101.6], "3 pages RGB 300x400 @100dpi")

    if large:
        W, H = 9933, 14043
        arr = np.full((H, W), 255, dtype=np.uint8)
        arr[::118, :] = 0
        arr[:, ::118] = 0
        Image.fromarray(arr, "L").save(out / "a0_300dpi.tif", compression="tiff_lzw", dpi=(300, 300))
        note("a0_300dpi.tif", "tiff", 1, [841, 1189], "8-bit gray LZW A0 @300dpi (--large)")


# ---------------------------------------------------------------- malicious
def png_chunk(t: bytes, d: bytes) -> bytes:
    return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)


def make_malicious(out: Path):
    m = out / "malicious"
    m.mkdir(parents=True, exist_ok=True)
    R = "malicious/"
    bomb = (b"\x89PNG\r\n\x1a\n" + png_chunk(b"IHDR", struct.pack(">IIBBBBB", 100000, 100000, 8, 2, 0, 0, 0))
            + png_chunk(b"IDAT", zlib.compress(b"\x00" * 64, 9)) + png_chunk(b"IEND", b""))
    (m / "pixel_bomb.png").write_bytes(bomb)
    note(R + "pixel_bomb.png", "png", 1, None, "header claims 100000x100000 px; tiny file", expect="reject")

    n = 10000
    objs = [b"<</Type/Catalog/Pages 2 0 R>>",
            b"<</Type/Pages/Count %d/Kids[%s]>>" % (n, b" ".join(b"%d 0 R" % (i + 3) for i in range(n)))]
    objs += [b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 10 10]>>"] * n
    (m / "pdf_10000_pages.pdf").write_bytes(raw_pdf(objs))
    note(R + "pdf_10000_pages.pdf", "pdf", n, [3.5, 3.5], "10000 tiny pages; must hit page limit", expect="reject")

    (m / "svg_xxe.svg").write_bytes(
        b'<?xml version="1.0"?>\n<!DOCTYPE svg [\n  <!ENTITY xxe SYSTEM "file:///etc/passwd">\n]>\n'
        b'<svg xmlns="http://www.w3.org/2000/svg" width="100mm" height="100mm">\n'
        b'  <text x="10" y="20">&xxe;</text>\n</svg>\n')
    note(R + "svg_xxe.svg", "svg", 1, [100, 100], "external entity file:///etc/passwd", expect="reject_or_neutralize")
    (m / "svg_remote_href.svg").write_bytes(
        b'<?xml version="1.0"?>\n<svg xmlns="http://www.w3.org/2000/svg" '
        b'xmlns:xlink="http://www.w3.org/1999/xlink" width="100mm" height="100mm">\n'
        b'  <image xlink:href="http://127.0.0.1:9/evil.png" x="0" y="0" width="50" height="50"/>\n'
        b'  <image href="file:///etc/passwd" x="0" y="50" width="50" height="50"/>\n</svg>\n')
    note(R + "svg_remote_href.svg", "svg", 1, [100, 100], "remote/file href; must not be fetched", expect="reject_or_neutralize")
    (m / "ps_read_passwd.ps").write_bytes(
        b"%!PS-Adobe-3.0\n%%BoundingBox: 0 0 300 300\n/Courier findfont 12 scalefont setfont\n"
        b"(/etc/passwd) (r) file 200 string readline pop 20 200 moveto show\n"
        b"(/tmp/a0td_pwned) (w) file dup (pwned) writestring closefile\nshowpage\n")
    note(R + "ps_read_passwd.ps", "ps", 1, None, "tries to read /etc/passwd and write /tmp (needs -dSAFER)", expect="reject_or_neutralize")

    objs = [b"<</Type/Catalog/Pages 2 0 R/Names<</JavaScript<</Names[(init) 5 0 R]>>>>/OpenAction 5 0 R>>",
            b"<</Type/Pages/Count 1/Kids[3 0 R]>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R/Resources<<>>>>",
            b"<</Length 0>>\nstream\n\nendstream",
            b"<</S/JavaScript/JS(app.alert\\('pwned'\\);)>>"]
    (m / "pdf_javascript.pdf").write_bytes(raw_pdf(objs))
    note(R + "pdf_javascript.pdf", "pdf", 1, [70.6, 70.6], "OpenAction JavaScript; must be neutralized (never executed)", expect="accept_neutralized")

    (m / "misleading_ext.pdf").write_bytes(
        b"\x89PNG\r\n\x1a\n" + png_chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + png_chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + png_chunk(b"IEND", b""))
    note(R + "misleading_ext.pdf", "png", 1, None, "PNG content with .pdf extension; detect by magic bytes", expect="detect_as_png")

    parts = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/><Override PartName="/word/document.xml" ContentType="application/vnd.ms-word.document.macroEnabled.main+xml"/></Types>',
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        "word/document.xml": '<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>macro placeholder</w:t></w:r></w:p></w:body></w:document>',
        "word/_rels/document.xml.rels": '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
        "word/vbaProject.bin": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"PLACEHOLDER-NOT-REAL-VBA\x00" * 8,
    }
    with zipfile.ZipFile(m / "docx_macro.docm", "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in parts.items():
            zi = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, content)
    note(R + "docx_macro.docm", "docm", 1, None, "macro-enabled docx with fake vbaProject.bin placeholder (not real VBA)", expect="reject")


# ---------------------------------------------------------------- F2 formats
def rezip(path: Path):
    """Rewrite a ZIP with fixed timestamps and member order (deterministic Office files)."""
    with zipfile.ZipFile(path) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in sorted(items):
            zi = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, data)


def make_office(out: Path):
    """pptx / docx through python-pptx / python-docx (dev-only dependencies; skipped if absent)."""
    import datetime as dt
    try:
        from docx import Document
        from docx.shared import Mm as DMm
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Mm, Pt
    except ImportError:
        print("python-pptx / python-docx missing: skipping pptx and docx fixtures")
        return
    fixed = dt.datetime(2020, 1, 1)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(841), Mm(1189)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(100), Mm(100), Mm(641), Mm(989))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xCC, 0xDD, 0xFF)
    box.line.color.rgb = RGBColor(0, 0, 0)
    box.line.width = Pt(6)
    tb = slide.shapes.add_textbox(Mm(150), Mm(200), Mm(541), Mm(200))
    tb.text_frame.text = "A0 slide"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(160)
    cp = prs.core_properties
    cp.created = cp.modified = fixed
    cp.author = cp.last_modified_by = ""
    cp.title = "fixture"
    prs.save(out / "pptx_a0.pptx")
    rezip(out / "pptx_a0.pptx")
    note("pptx_a0.pptx", "pptx", 1, [841, 1189], "python-pptx, single A0 portrait slide")

    for name, font, text in (
        ("docx_simple.docx", None, "Plain A4 document"),
        ("docx_missing_font.docx", "Zapfino Imaginary", "Uses a font that is not installed"),
    ):
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = DMm(210), DMm(297)
        doc.add_heading("a0td fixture", 1)
        run = doc.add_paragraph().add_run(text)
        if font:
            run.font.name = font
        cp = doc.core_properties
        cp.created = cp.modified = cp.last_printed = fixed
        cp.author = cp.last_modified_by = ""
        cp.title = "fixture"
        doc.save(out / name)
        rezip(out / name)
    note("docx_simple.docx", "docx", 1, [210, 297], "python-docx, A4, default fonts")
    note("docx_missing_font.docx", "docx", 1, [210, 297],
         "run in font 'Zapfino Imaginary' (not installed): FONTS_SUBSTITUTED")


def make_svgs(out: Path):
    body = ('<rect x="5" y="5" width="190" height="90" fill="#ccddff" stroke="#000" stroke-width="1"/>'
            '<circle cx="100" cy="50" r="30" fill="none" stroke="#c00" stroke-width="2"/>'
            '<text x="12" y="20" font-family="sans-serif" font-size="10">svg fixture</text>')
    head = '<?xml version="1.0" encoding="UTF-8"?>\n<svg xmlns="http://www.w3.org/2000/svg" '
    (out / "svg_physical.svg").write_text(
        head + 'width="200mm" height="100mm" viewBox="0 0 200 100">' + body + '</svg>\n')
    note("svg_physical.svg", "svg", 1, [200, 100], "width/height in mm")
    (out / "svg_no_units.svg").write_text(head + 'viewBox="0 0 200 100">' + body + '</svg>\n')
    note("svg_no_units.svg", "svg", 1, None,
         "viewBox only, no physical units: PHYSICAL_SIZE_ASSUMED, fits A0 landscape (1189x594.5)")


def make_postscript(out: Path):
    body = ("0.5 setlinewidth newpath 20 20 moveto 575 20 lineto 575 822 lineto 20 822 lineto "
            "closepath stroke\n1 0 0 setrgbcolor newpath 20 20 moveto 575 822 lineto stroke\n"
            "0 setgray /Helvetica findfont 40 scalefont setfont 60 400 moveto (PS fixture) show\n")
    (out / "eps_a4.eps").write_text(
        "%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 595 842\n%%Pages: 1\n%%EndComments\n"
        + body + "showpage\n%%EOF\n")
    note("eps_a4.eps", "eps", 1, [209.9, 297], "hand-written EPS, BoundingBox 595x842 pt")
    (out / "ps_a4.ps").write_text(
        "%!PS-Adobe-3.0\n%%Pages: 1\n%%EndComments\n<< /PageSize [595 842] >> setpagedevice\n"
        + body + "showpage\n%%EOF\n")
    note("ps_a4.ps", "ps", 1, [209.9, 297], "hand-written PostScript, setpagedevice A4")
    (out / "ai_postscript.ai").write_text(
        "%!PS-Adobe-3.0 EPSF-3.0\n%%Creator: Adobe Illustrator(R) 8.0\n%%BoundingBox: 0 0 595 842\n"
        "%%EndComments\n" + body + "showpage\n%%EOF\n")
    note("ai_postscript.ai", "ai", 1, [209.9, 297], "PostScript-based Illustrator file -> Ghostscript")
    grid_pdf(out / "ai_pdf.ai", 297, 420)
    note("ai_pdf.ai", "ai", 1, [297, 420], "PDF-compatible Illustrator file -> PDF path")


def make_xps(out: Path):
    ns = "http://schemas.microsoft.com/xps/2005/06"
    hdr = '<?xml version="1.0" encoding="utf-8"?>'
    parts = {
        "[Content_Types].xml": (
            hdr + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="fdseq" ContentType="application/vnd.ms-package.xps-fixeddocumentsequence+xml"/>'
            '<Default Extension="fdoc" ContentType="application/vnd.ms-package.xps-fixeddocument+xml"/>'
            '<Default Extension="fpage" ContentType="application/vnd.ms-package.xps-fixedpage+xml"/></Types>'),
        "_rels/.rels": (
            hdr + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="R0" Type="http://schemas.microsoft.com/xps/2005/06/fixedrepresentation" '
            'Target="/FixedDocumentSequence.fdseq"/></Relationships>'),
        "FixedDocumentSequence.fdseq": (
            hdr + f'<FixedDocumentSequence xmlns="{ns}">'
            '<DocumentReference Source="/Documents/1/FixedDocument.fdoc"/></FixedDocumentSequence>'),
        "Documents/1/FixedDocument.fdoc": (
            hdr + f'<FixedDocument xmlns="{ns}"><PageContent Source="Pages/1.fpage"/></FixedDocument>'),
        "Documents/1/Pages/1.fpage": (
            hdr + f'<FixedPage xmlns="{ns}" Width="1122.52" Height="793.7" xml:lang="en-US">'
            '<Path Data="M 50,50 L 1070,50 L 1070,740 L 50,740 Z" Fill="#FFCCDDFF" '
            'Stroke="#FF000000" StrokeThickness="4"/>'
            '<Path Data="M 50,50 L 1070,740" Stroke="#FFCC0000" StrokeThickness="3"/></FixedPage>'),
    }
    with zipfile.ZipFile(out / "xps_a4.xps", "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in parts.items():
            zi = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, content)
    note("xps_a4.xps", "xps", 1, [297, 210], "hand-written XPS, A4 landscape, two paths")


def make_modern_rasters(out: Path):
    w, h = 800, 600
    yy, xx = np.mgrid[0:h, 0:w]
    arr = np.stack([xx * 255 // w, yy * 255 // h, (xx + yy) * 255 // (w + h)], axis=-1)
    arr = arr.astype(np.uint8)
    arr[::50, :] = 0
    arr[:, ::50] = 0
    img = Image.fromarray(arr, "RGB")
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        exif = Image.Exif()
        exif[282], exif[283], exif[296] = 150.0, 150.0, 2
        img.save(out / "photo_150dpi.heic", format="HEIF", quality=85, exif=exif.tobytes())
        note("photo_150dpi.heic", "heif", 1, [135.5, 101.6], "800x600 HEIC with EXIF 150 dpi")
        getattr(pillow_heif, "register_avif_opener", lambda: None)()
        img.save(out / "photo_nodpi.avif", format="AVIF", quality=85)
        note("photo_nodpi.avif", "avif", 1, None, "800x600 AVIF, no DPI: PHYSICAL_SIZE_ASSUMED")
    except Exception as exc:
        print("HEIC / AVIF encoding unavailable:", exc)
    try:
        img.save(out / "photo_nodpi.jp2", format="JPEG2000")
        note("photo_nodpi.jp2", "jp2", 1, None, "800x600 JPEG 2000, no DPI: PHYSICAL_SIZE_ASSUMED")
    except Exception as exc:
        print("JPEG 2000 encoding unavailable:", exc)


def make_metafiles(out: Path):
    """EMF / WMF through LibreOffice (only when soffice exists; skipped otherwise)."""
    import shutil
    import subprocess
    import tempfile
    if not shutil.which("soffice"):
        print("soffice missing: skipping emf / wmf fixtures")
        return
    src = out / "svg_physical.svg"
    for ext in ("emf", "wmf"):
        with tempfile.TemporaryDirectory() as tmp:
            try:
                subprocess.run(
                    ["soffice", "--headless", "--norestore",
                     f"-env:UserInstallation=file://{tmp}/profile",
                     "--convert-to", ext, "--outdir", tmp, str(src)],
                    check=True, capture_output=True, timeout=180,
                    env={"HOME": tmp, "PATH": "/usr/bin:/bin:/usr/local/bin"})
                data = (Path(tmp) / f"svg_physical.{ext}").read_bytes()
            except Exception as exc:
                print(f"{ext} export failed:", exc)
                continue
        (out / f"metafile_a4.{ext}").write_bytes(data)
        note(f"metafile_a4.{ext}", ext, 1, None, f"{ext.upper()} exported by LibreOffice from an SVG")


def make_f2_malicious(out: Path):
    """ODT whose document-open event runs a Basic macro writing /tmp/a0td_macro_pwned."""
    m = out / "malicious"
    m.mkdir(parents=True, exist_ok=True)
    hdr = '<?xml version="1.0" encoding="UTF-8"?>'
    mt = "application/vnd.oasis.opendocument.text"
    entries = "".join(
        f'<manifest:file-entry manifest:full-path="{p}" manifest:media-type="{t}"/>'
        for p, t in (("/", mt), ("content.xml", "text/xml"), ("Basic/script-lc.xml", "text/xml"),
                     ("Basic/Standard/script-lb.xml", "text/xml"),
                     ("Basic/Standard/Module1.xml", "text/xml")))
    manifest = (hdr + '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:'
                'manifest:1.0" manifest:version="1.2">' + entries + '</manifest:manifest>')
    href = ('vnd.sun.star.script:Standard.Module1.Main?language=Basic&amp;location=document')
    listeners = "".join(
        f'<script:event-listener script:language="ooo:script" script:event-name="{ev}" '
        f'xlink:href="{href}" xlink:type="simple"/>' for ev in ("dom:load", "office:load"))
    content = (
        hdr + '<office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:script="urn:oasis:names:tc:opendocument:xmlns:script:1.0" '
        'xmlns:xlink="http://www.w3.org/1999/xlink" office:version="1.2">'
        '<office:scripts><office:event-listeners>' + listeners + '</office:event-listeners>'
        '</office:scripts><office:body><office:text><text:p>macro document</text:p>'
        '</office:text></office:body></office:document-content>')
    lc = (hdr + '<library:libraries xmlns:library="http://openoffice.org/2000/library" '
          'xmlns:xlink="http://www.w3.org/1999/xlink"><library:library library:name="Standard" '
          'xlink:href="$(USER)/basic/Standard/script.xlb/" xlink:type="simple" '
          'library:link="false"/></library:libraries>')
    lb = (hdr + '<library:library xmlns:library="http://openoffice.org/2000/library" '
          'library:name="Standard" library:readonly="false" library:passwordprotected="false">'
          '<library:element library:name="Module1"/></library:library>')
    mod = (hdr + '<script:module xmlns:script="http://openoffice.org/2000/script" '
           'script:name="Module1" script:language="StarBasic">'
           'Sub Main&#10;  Dim n As Integer&#10;  n = FreeFile()&#10;'
           '  Open &quot;/tmp/a0td_macro_pwned&quot; For Output As #n&#10;'
           '  Print #n, &quot;pwned&quot;&#10;  Close #n&#10;End Sub&#10;</script:module>')
    with zipfile.ZipFile(m / "odt_macro.odt", "w") as z:
        for name, data, comp in (
            ("mimetype", mt, zipfile.ZIP_STORED),
            ("META-INF/manifest.xml", manifest, zipfile.ZIP_DEFLATED),
            ("content.xml", content, zipfile.ZIP_DEFLATED),
            ("Basic/script-lc.xml", lc, zipfile.ZIP_DEFLATED),
            ("Basic/Standard/script-lb.xml", lb, zipfile.ZIP_DEFLATED),
            ("Basic/Standard/Module1.xml", mod, zipfile.ZIP_DEFLATED),
        ):
            zi = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            zi.compress_type = comp
            z.writestr(zi, data)
    note("malicious/odt_macro.odt", "odt", 1, None,
         "ODT whose open event runs a macro writing /tmp/a0td_macro_pwned; macros must stay off",
         expect="accept_neutralized")


# ---------------------------------------------------------------- F3 formats (CAD, HPGL, DjVu)
FIXED_GUID = "{00000000-0000-0000-0000-00000000A0D0}"


def _cad_doc():
    """A DXF R2010 in millimetres with fixed header timestamps / GUIDs (deterministic output)."""
    import ezdxf
    doc = ezdxf.new("R2010", setup=True)
    doc.units = ezdxf.units.MM
    hdr = doc.header
    for var in ("$TDCREATE", "$TDUCREATE", "$TDUPDATE", "$TDUUPDATE"):
        hdr[var] = 2458849.0
    hdr["$FINGERPRINTGUID"] = hdr["$VERSIONGUID"] = FIXED_GUID
    return doc


def _save_dxf(doc, path: Path):
    """Save and pin the two header values ezdxf rewrites on every save ($TDUPDATE, $VERSIONGUID)."""
    doc.saveas(path)
    data = path.read_bytes()
    nl = rb"[^\r\n]*"
    data = re.sub(rb"(\$TDUPDATE\r?\n\s*40\r?\n)" + nl, rb"\g<1>2458849.0", data)
    data = re.sub(rb"(\$VERSIONGUID\r?\n\s*2\r?\n)" + nl,
                  rb"\g<1>" + FIXED_GUID.encode(), data)
    data = re.sub(rb"(\d+\.\d+\.\d+ @ )\d{4}-\d\d-\d\dT[\d:.+]+",
                  rb"\g<1>2020-01-01T00:00:00+00:00", data)
    path.write_bytes(data)


def _draw_plan(msp):
    """A 12 x 8 m floor plan: walls, door arcs, round columns, hatched room, text, dimension."""
    msp.add_lwpolyline([(0, 0), (12000, 0), (12000, 8000), (0, 8000)], close=True,
                       dxfattribs={"layer": "WALLS", "color": 1, "lineweight": 50})
    msp.add_lwpolyline([(6000, 0), (6000, 5000)], dxfattribs={"layer": "WALLS", "color": 1, "lineweight": 35})
    msp.add_lwpolyline([(0, 5000), (9000, 5000)], dxfattribs={"layer": "WALLS", "color": 1, "lineweight": 35})
    msp.add_arc((2000, 5000), 900, 0, 90, dxfattribs={"layer": "DOORS", "color": 3})
    msp.add_arc((6000, 2000), 900, 90, 180, dxfattribs={"layer": "DOORS", "color": 3})
    for x in (3000, 9000):
        msp.add_circle((x, 6500), 250, dxfattribs={"layer": "COLUMNS", "color": 5})
    hatch = msp.add_hatch(color=8, dxfattribs={"layer": "ROOMS"})
    hatch.set_pattern_fill("ANSI31", scale=40)
    hatch.paths.add_polyline_path([(6200, 200), (11800, 200), (11800, 4800), (6200, 4800)])
    msp.add_text("LIVING ROOM", height=300, dxfattribs={"layer": "TEXT", "color": 7, "insert": (6500, 2200)})
    msp.add_text("12000 x 8000 mm plan", height=400, dxfattribs={"layer": "TEXT", "color": 2, "insert": (500, 7300)})
    dim = msp.add_linear_dim(base=(0, -1200), p1=(0, 0), p2=(12000, 0), dimstyle="EZDXF")
    dim.render()


def make_cad(out: Path):
    import shutil
    import subprocess
    import tempfile

    import ezdxf
    from ezdxf.enums import TextEntityAlignment

    # Model space only: unitless to scale (auto 1:20 on an A0 sheet at most).
    doc = _cad_doc()
    _draw_plan(doc.modelspace())
    _save_dxf(doc, out / "dxf_model_plan.dxf")
    note("dxf_model_plan.dxf", "dxf", 1, None,
         "12000x8000 mm plan in model space ($INSUNITS=4); no paper layout; auto scale 1:20 "
         "(page 615x415 mm); explicit 1:50 gives 255x175 mm")

    # Paper space: an A0 landscape layout with a frame, title block and one viewport.
    doc = _cad_doc()
    _draw_plan(doc.modelspace())
    doc.layouts.rename("Layout1", "A0 sheet")
    psp = doc.layouts.get("A0 sheet")
    psp.page_setup(size=(1189, 841), margins=(0, 0, 0, 0), units="mm", scale=(1, 1))
    psp.add_lwpolyline([(10, 10), (1179, 10), (1179, 831), (10, 831)], close=True,
                       dxfattribs={"lineweight": 50, "color": 7})
    psp.add_lwpolyline([(889, 10), (889, 80), (1179, 80)], dxfattribs={"lineweight": 25, "color": 7})
    psp.add_text("A0 layout fixture", height=12,
                 dxfattribs={"color": 7}).set_placement((900, 40), align=TextEntityAlignment.LEFT)
    vp = psp.add_viewport(center=(594.5, 450), size=(1100, 700), view_center_point=(6000, 4000),
                          view_height=7000 * 700 / 700 * 1.0)
    vp.dxf.status = 2
    _save_dxf(doc, out / "dxf_a0_layout.dxf")
    note("dxf_a0_layout.dxf", "dxf", 1, [1189, 841],
         "paper space layout 'A0 sheet' (1189x841 mm) with frame, title block and a viewport on the "
         "plan; layouts: Model, 'A0 sheet'")

    # DWG: LibreDWG's dxf2dwg on the plan (skipped when it is not installed).
    dxf2dwg = shutil.which("dxf2dwg") or ("/opt/libredwg/bin/dxf2dwg"
                                          if Path("/opt/libredwg/bin/dxf2dwg").exists() else None)
    if dxf2dwg is None:
        print("dxf2dwg missing: skipping the DWG fixture")
        return
    with tempfile.TemporaryDirectory() as tmp:
        target = Path(tmp) / "plan.dwg"
        try:
            subprocess.run([dxf2dwg, "-y", "--as", "r2000", "-o", str(target),
                            str(out / "dxf_model_plan.dxf")],
                           check=True, capture_output=True, timeout=120)
            data = target.read_bytes()
        except Exception as exc:
            print("dxf2dwg failed, skipping the DWG fixture:", exc)
            return
    (out / "dwg_model_plan.dwg").write_bytes(data)
    note("dwg_model_plan.dwg", "dwg", 1, None,
         "dxf_model_plan.dxf converted to DWG R2000 with LibreDWG dxf2dwg")


def make_hpgl(out: Path):
    """A text HPGL file: A1 landscape (33640 x 23760 plotter units of 0.025 mm) with a frame, a
    10 cm grid, a circle and a diagonal (pen 1)."""
    lines = ["IN;", "SP1;", "PU0,0;PD33640,0,33640,23760,0,23760,0,0;"]
    for x in range(4000, 33640, 4000):
        lines.append(f"PU{x},0;PD{x},23760;")
    for y in range(4000, 23760, 4000):
        lines.append(f"PU0,{y};PD33640,{y};")
    lines += ["SP2;", "PU16820,11880;CI6000;", "PU0,0;PD33640,23760;", "PU0,0;SP0;"]
    text = "\n".join(lines) + "\n"
    (out / "hpgl_plot.plt").write_bytes(text.encode("ascii"))
    note("hpgl_plot.plt", "hpgl", 1, [843.8, 596.6],
         "HPGL plot of 841x594 mm; the frame is sized from the pen coordinates plus 0.1 in")


def make_djvu(out: Path):
    """A4 colour page at 150 dpi through ``c44`` (DjVuLibre, skipped when it is not installed)."""
    import shutil
    import subprocess
    import tempfile
    if shutil.which("c44") is None:
        print("c44 missing: skipping the DjVu fixture")
        return
    w, h = 1240, 1754
    yy, xx = np.mgrid[0:h, 0:w]
    arr = np.stack([xx * 255 // w, yy * 255 // h, (xx + yy) * 255 // (w + h)], axis=-1).astype(np.uint8)
    arr[::100, :] = 0
    arr[:, ::100] = 0
    with tempfile.TemporaryDirectory() as tmp:
        ppm = Path(tmp) / "scan.ppm"
        ppm.write_bytes(b"P6\n%d %d\n255\n" % (w, h) + arr.tobytes())
        target = Path(tmp) / "scan.djvu"
        try:
            subprocess.run(["c44", "-dpi", "150", str(ppm), str(target)],
                           check=True, capture_output=True, timeout=120)
            data = target.read_bytes()
        except Exception as exc:
            print("c44 failed, skipping the DjVu fixture:", exc)
            return
    (out / "djvu_a4_150dpi.djvu").write_bytes(data)
    note("djvu_a4_150dpi.djvu", "djvu", 1, [210, 297], "1240x1754 colour page at 150 dpi (A4)")


def _fixed_hash_seed():
    """ezdxf orders some DXF classes through a set: pin the hash seed so the output is stable."""
    import os
    import sys
    if os.environ.get("PYTHONHASHSEED") != "0":
        os.environ["PYTHONHASHSEED"] = "0"
        os.execv(sys.executable, [sys.executable, *sys.argv])


def main():
    _fixed_hash_seed()
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="fixtures")
    ap.add_argument("--large", action="store_true", help="also generate a0_300dpi.tif")
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    make_pdfs(out)
    make_rasters(out, args.large)
    make_malicious(out)
    make_office(out)
    make_svgs(out)
    make_postscript(out)
    make_xps(out)
    make_modern_rasters(out)
    make_metafiles(out)
    make_f2_malicious(out)
    make_cad(out)
    make_hpgl(out)
    make_djvu(out)
    for rel in MANIFEST:
        MANIFEST[rel]["bytes"] = (out / rel).stat().st_size
    (out / "manifest.json").write_bytes((json.dumps(MANIFEST, indent=2, sort_keys=True) + "\n").encode())
    for rel in sorted(MANIFEST):
        print(hashlib.sha256((out / rel).read_bytes()).hexdigest(), rel)


if __name__ == "__main__":
    main()
